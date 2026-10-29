from pathlib import Path

import openpyxl

from tenderai.align import align, canon_key
from tenderai.inventory import detect_format
from tenderai.models import Requirement, Span
from tenderai.parse import parse_constraint
from tenderai.pipeline import run


def test_magic_bytes_ignore_the_extension():
    assert detect_format(b"PK\x03\x04" + b"word/document.xml") == "docx"
    assert detect_format(b"%PDF-1.4") == "pdf"
    assert detect_format(b'{"error":"missing"}') == "error_json"
    assert detect_format(b"Rar!\x1a\x07") == "rar"
    assert detect_format(b"\xd0\xcf\x11\xe0" + "Workbook".encode("utf-16le")) == "xls"
    assert detect_format(b"\xd0\xcf\x11\xe0" + "WordDocument".encode("utf-16le")) == "doc"


def test_bolus_lock_does_not_steal_bolus():
    assert canon_key("Система блокировки болюса") == "bolus_lock"
    assert canon_key("Болюс") == "bolus"
    assert canon_key("Номинальный объем резервуара") == "reservoir_volume"


def test_pipeline_matches_a_single_position(tmp_path: Path):
    dataset = tmp_path / "dataset"
    _write_sku(dataset / "SKU" / "CT0050L.xlsx")
    _write_tender(dataset / "tenders" / "01_0001" / "ooz.xlsx")
    summary = run(dataset, tmp_path / "out")
    assert summary["requirements"] == 2
    row = summary["tender_rows"][0]
    assert row["sku_id"] == "CT0050L"
    assert row["status"] == "match"


def test_unaligned_requirement_is_unknown_not_mismatch(tmp_path: Path):
    dataset = tmp_path / "dataset"
    _write_sku(dataset / "SKU" / "CT0050L.xlsx")
    path = dataset / "tenders" / "01_0001" / "ooz.xlsx"
    path.parent.mkdir(parents=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["E1"] = "Наименование характеристики"
    sheet["F1"] = "Значение характеристики"
    sheet["B1"] = "Наименование товара"
    sheet["B2"] = "Насос инфузионный эластомерный"
    sheet["E2"] = "Цвет транспортной тары"
    sheet["F2"] = "Синий"
    book.save(path)
    book.close()
    summary = run(dataset, tmp_path / "out")
    assert summary["tender_rows"][0]["status"] == "unknown"
    assert summary["tender_rows"][0]["reason"] == "no_confirmed_position"


def test_docx_table_requirements_use_pokazatel_headers(tmp_path: Path):
    import docx

    dataset = tmp_path / "dataset"
    _write_sku(dataset / "SKU" / "CT0050L.xlsx")
    path = dataset / "tenders" / "01_0001" / "opis.docx"
    path.parent.mkdir(parents=True)
    document = docx.Document()
    table = document.add_table(rows=3, cols=5)
    header = ["№", "Наименование", "Наименование показателя", "Значение показателя", "Единица измерения показателя"]
    for col, text in enumerate(header):
        table.cell(0, col).text = text
    row1 = ["1", "Насос инфузионный", "Номинальный объем резервуара", "≥ 200 и ≤ 300", "мл"]
    row2 = ["1", "Насос инфузионный", "Болюс", "Нет", ""]
    for col, text in enumerate(row1):
        table.cell(1, col).text = text
    for col, text in enumerate(row2):
        table.cell(2, col).text = text
    document.save(path)
    summary = run(dataset, tmp_path / "out")
    assert summary["requirements"] == 2
    assert summary["tender_rows"][0]["status"] == "match"


def test_alignment_uses_canon_before_tokens():
    requirement = Requirement(
        requirement_id="r",
        tender_id="1",
        position_id="p",
        position_title="Насос",
        name="Номинальный объем резервуара",
        raw_text="275",
        mandatory=True,
        constraint=parse_constraint("275", unit="мл"),
        span=Span("f", "A1", "275"),
    )
    characteristics = load_skus_from_memory()
    chosen, method, score = align(requirement, characteristics)
    assert chosen is not None
    assert chosen.name.startswith("Номинальный")
    assert method == "canon"
    assert score == 1.0


def load_skus_from_memory():
    from tenderai.models import SkuCharacteristic

    return [
        SkuCharacteristic(
            sku_id="CT0050L",
            characteristic_id="CT0050L:1",
            name="Номинальный объём резервуара",
            value_raw="250",
            unit_raw="мл",
            knowledge="stated",
            constraint=parse_constraint("250", unit="мл"),
            span=Span("sku", "A1", "250"),
        )
    ]


def _write_sku(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = "CT0050L"
    sheet["B1"] = "Accufuser"
    headers = ["Характеристика", "Значение", "Единица", "Сведения", "", "Источник", "Формулировка"]
    for index, header in enumerate(headers, start=1):
        if header:
            sheet.cell(2, index, header)
    rows = [
        ("Номинальный объём резервуара", "250", "мл", "Указано", "250"),
        ("Болюс", "Нет", "", "Указано", "Нет"),
    ]
    for offset, row in enumerate(rows, start=3):
        sheet.cell(offset, 1, row[0])
        sheet.cell(offset, 2, row[1])
        sheet.cell(offset, 3, row[2])
        sheet.cell(offset, 4, row[3])
        sheet.cell(offset, 7, row[4])
    book.save(path)
    book.close()


def _write_tender(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["B1"] = "Наименование товара"
    sheet["E1"] = "Наименование характеристики"
    sheet["F1"] = "Значение характеристики"
    sheet["M1"] = "Единица измерения характеристики"
    sheet["B2"] = "Насос инфузионный эластомерный"
    sheet["E2"] = "Номинальный объем резервуара"
    sheet["F2"] = "≥ 200 и ≤ 300"
    sheet["M2"] = "мл"
    sheet["E3"] = "Болюс"
    sheet["F3"] = "Нет"
    book.save(path)
    book.close()
