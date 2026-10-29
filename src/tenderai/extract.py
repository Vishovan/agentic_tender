"""Pull text and tables out of a file. OCR is intentionally absent."""

from __future__ import annotations

import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from tenderai.inventory import FileRecord


@dataclass
class Cell:
    a1: str
    value: str


@dataclass
class Table:
    sheet: str
    rows: list[list[Cell]]


@dataclass
class Document:
    record: FileRecord
    status: str
    error: str | None = None
    text: str = ""
    tables: list[Table] = field(default_factory=list)


def extract_document(dataset: Path, record: FileRecord) -> Document:
    path = dataset / record.path
    try:
        if record.detected_format == "xlsx":
            tables = _xlsx_tables(path)
            return Document(record, "ok", tables=tables, text=_tables_text(tables))
        if record.detected_format == "xls":
            tables = _xls_tables(path)
            return Document(record, "ok", tables=tables, text=_tables_text(tables))
        if record.detected_format == "docx":
            text, tables = _docx(path)
            return Document(record, "ok", text=text, tables=tables)
        if record.detected_format == "pdf":
            text = _pdf(path)
            status = "ok" if text.strip() else "empty_text"
            return Document(record, status, text=text, error=None if text.strip() else "pdf_text_layer_empty")
        if record.detected_format == "xml":
            return Document(record, "ok", text=_xml(path))
        if record.detected_format == "txt":
            return Document(record, "ok", text=path.read_text(encoding="utf-8", errors="replace"))
        if record.detected_format == "doc":
            text = _doc(path)
            return Document(record, "ok", text=text)
        if record.detected_format == "error_json":
            return Document(record, "gap", error="saved_error_response", text=path.read_text(encoding="utf-8", errors="replace"))
        if record.detected_format == "rar":
            return Document(record, "gap", error="rar_not_unpacked")
        return Document(record, "gap", error=f"unsupported:{record.detected_format}")
    except Exception as exc:  # a broken file must not drop the tender
        return Document(record, "error", error=f"{type(exc).__name__}: {exc}")


def _tables_text(tables: list[Table]) -> str:
    lines: list[str] = []
    for table in tables:
        for row in table.rows:
            lines.append(" | ".join(cell.value for cell in row if cell.value))
    return "\n".join(lines)


def _xlsx_tables(path: Path) -> list[Table]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    tables: list[Table] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[Cell]] = []
            for row in sheet.iter_rows():
                cells = [
                    Cell(a1=cell.coordinate, value=_text(cell.value))
                    for cell in row
                    if _text(cell.value)
                ]
                if cells:
                    rows.append(cells)
            if rows:
                tables.append(Table(sheet=sheet.title, rows=rows))
    finally:
        workbook.close()
    return tables


def _xls_tables(path: Path) -> list[Table]:
    import xlrd

    book = xlrd.open_workbook(path)
    tables: list[Table] = []
    for sheet in book.sheets():
        rows: list[list[Cell]] = []
        for row_index in range(sheet.nrows):
            cells: list[Cell] = []
            for col_index in range(sheet.ncols):
                value = _text(sheet.cell_value(row_index, col_index))
                if not value:
                    continue
                cells.append(Cell(a1=_a1(row_index + 1, col_index + 1), value=value))
            if cells:
                rows.append(cells)
        if rows:
            tables.append(Table(sheet=sheet.name, rows=rows))
    return tables


def _docx(path: Path) -> tuple[str, list[Table]]:
    import docx

    document = docx.Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[Table] = []
    for index, table in enumerate(document.tables, start=1):
        rows: list[list[Cell]] = []
        for row_index, row in enumerate(table.rows, start=1):
            cells = []
            for col_index, cell in enumerate(row.cells, start=1):
                value = " ".join(cell.text.split())
                if value:
                    cells.append(Cell(a1=f"T{index}R{row_index}C{col_index}", value=value))
            if cells:
                rows.append(cells)
        if rows:
            tables.append(Table(sheet=f"table-{index}", rows=rows))
    return "\n".join(paragraphs + [_tables_text(tables)]), tables


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(f"[page {index}]\n{text}" for index, text in enumerate(pages, start=1) if text)


def _xml(path: Path) -> str:
    from xml.etree import ElementTree

    tree = ElementTree.parse(path)
    return " ".join(piece.strip() for piece in tree.getroot().itertext() if piece.strip())


def _doc(path: Path) -> str:
    completed = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        check=False,
        capture_output=True,
    )
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.decode("utf-8", errors="replace")[:500] or "textutil failed")
    return completed.stdout.decode("utf-8", errors="replace")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\u00a0", " ").strip()


def _a1(row: int, col: int) -> str:
    letters = ""
    index = col
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"
